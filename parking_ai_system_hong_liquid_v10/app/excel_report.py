"""
QL-08: Xuất báo cáo lưu lượng/doanh thu ra file Excel (.xlsx), dùng cùng
số liệu đã tổng hợp cho AI (app/ai_data.py) để đảm bảo khớp với những gì
hiển thị trên màn hình Thống kê - không tính lại theo công thức khác.
"""
from datetime import datetime
from io import BytesIO
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app import ai_data, crud, schemas

_MAU_HEADER = "A71350"   # trùng --pri-dark của giao diện web
_MAU_HEADER_CHU = "FFFFFF"
_MAU_TIEU_DE = "D92D73"  # trùng --pri

NHAN_HINH_THUC_GUI = {
    "sang": "Sáng (6h-13h)", "chieu": "Chiều (13h-18h)", "toi": "Tối (18h-6h)",
    "qua_dem": "Gửi qua đêm", "thang": "Gửi theo tháng", "qua_dem_thang": "Gửi xe qua đêm theo tháng",
}


def _ke_hang_tieu_de(ws, dong: int, tieu_de: str, so_cot: int):
    ws.merge_cells(start_row=dong, start_column=1, end_row=dong, end_column=so_cot)
    o = ws.cell(row=dong, column=1, value=tieu_de)
    o.font = Font(bold=True, size=13, color=_MAU_TIEU_DE)
    o.alignment = Alignment(horizontal="left")


def _ke_hang_header(ws, dong: int, cot_bat_dau: int, nhan: list[str]):
    for i, ten in enumerate(nhan):
        o = ws.cell(row=dong, column=cot_bat_dau + i, value=ten)
        o.font = Font(bold=True, color=_MAU_HEADER_CHU)
        o.fill = PatternFill("solid", fgColor=_MAU_HEADER)
        o.alignment = Alignment(horizontal="center", vertical="center")


def _tu_dong_gian_cot(ws, do_rong: dict[str, int]):
    for cot, rong in do_rong.items():
        ws.column_dimensions[cot].width = rong


def tao_bao_cao_excel(db: Session, tu_ngay: datetime, den_ngay: datetime,
                       khu_vuc_id: Optional[int] = None) -> BytesIO:
    ll = ai_data.thong_ke_luu_luong(db, tu_ngay, den_ngay, khu_vuc_id)
    dt = ai_data.thong_ke_doanh_thu(db, tu_ngay, den_ngay, khu_vuc_id)
    ld = ai_data.ty_le_lap_day(db, khu_vuc_id)

    wb = Workbook()

    # ----- Sheet 1: Tổng quan -----
    ws = wb.active
    ws.title = "Tổng quan"
    _ke_hang_tieu_de(ws, 1, "BÁO CÁO VẬN HÀNH BÃI ĐỖ XE", 2)
    ws.cell(row=2, column=1, value="Khoảng thời gian").font = Font(bold=True)
    ws.cell(row=2, column=2, value=f"{tu_ngay:%d/%m/%Y %H:%M} - {den_ngay:%d/%m/%Y %H:%M}")
    ws.cell(row=3, column=1, value="Ngày xuất báo cáo").font = Font(bold=True)
    ws.cell(row=3, column=2, value=datetime.now().strftime("%d/%m/%Y %H:%M"))

    thoi_gian_dong_nhat = ll.get("thoi_gian_dong_nhat")
    nhan_dong_nhat = NHAN_HINH_THUC_GUI.get(thoi_gian_dong_nhat, thoi_gian_dong_nhat or "-")

    hang = 5
    for nhan, gia_tri in [
        ("Tổng số lượt xe", ll["tong_so_luot"]),
        ("Thời gian gửi xe đông nhất", nhan_dong_nhat),
        ("Tổng doanh thu (VNĐ)", dt["tong_doanh_thu"]),
    ]:
        ws.cell(row=hang, column=1, value=nhan).font = Font(bold=True)
        ws.cell(row=hang, column=2, value=gia_tri)
        hang += 1

    hang += 1
    _ke_hang_tieu_de(ws, hang, "Tỷ lệ lấp đầy theo khu vực (thời điểm xuất báo cáo)", 4)
    hang += 1
    _ke_hang_header(ws, hang, 1, ["Khu vực", "Tổng vị trí", "Đã đỗ", "Tỷ lệ lấp đầy"])
    for kv in ld["theo_khu_vuc"]:
        hang += 1
        ws.cell(row=hang, column=1, value=kv["khu_vuc"])
        ws.cell(row=hang, column=2, value=kv["tong_vi_tri"])
        ws.cell(row=hang, column=3, value=kv["da_do"])
        ws.cell(row=hang, column=4, value=f"{kv['ty_le']*100:.0f}%")
    _tu_dong_gian_cot(ws, {"A": 30, "B": 22, "C": 12, "D": 14})

    # ----- Sheet 2: Theo ngày -----
    ws2 = wb.create_sheet("Theo ngày")
    _ke_hang_header(ws2, 1, 1, ["Ngày", "Số lượt vào", "Doanh thu (VNĐ)"])
    tat_ca_ngay = sorted(set(ll["theo_ngay"]) | set(dt["theo_ngay"]))
    for i, ngay in enumerate(tat_ca_ngay, start=2):
        ws2.cell(row=i, column=1, value=ngay)
        ws2.cell(row=i, column=2, value=ll["theo_ngay"].get(ngay, {}).get("vao", 0))
        ws2.cell(row=i, column=3, value=dt["theo_ngay"].get(ngay, 0))
    _tu_dong_gian_cot(ws2, {"A": 16, "B": 16, "C": 18})

    # ----- Sheet 3: Theo thời gian gửi xe -----
    ws3 = wb.create_sheet("Theo thời gian gửi xe")
    _ke_hang_header(ws3, 1, 1, ["Thời gian gửi xe", "Số lượt", "Tỷ lệ %"])
    for i, (khoa, nhan) in enumerate(NHAN_HINH_THUC_GUI.items(), start=2):
        ws3.cell(row=i, column=1, value=nhan)
        ws3.cell(row=i, column=2, value=ll["theo_hinh_thuc_gui"].get(khoa, 0))
        ws3.cell(row=i, column=3, value=f"{ll['ty_le_theo_hinh_thuc_gui'].get(khoa, 0)}%")
    _tu_dong_gian_cot(ws3, {"A": 30, "B": 12, "C": 12})

    # ----- Sheet 4: Chi tiết lượt gửi xe -----
    ws4 = wb.create_sheet("Chi tiết lượt gửi xe")
    cot4 = ["Biển số", "Chủ xe", "SĐT", "Loại xe", "Khu vực", "Vị trí",
            "Thời gian gửi", "Giờ vào", "Giờ ra", "Phí thu (VNĐ)", "Mất vé", "Trạng thái"]
    _ke_hang_header(ws4, 1, 1, cot4)
    chi_tiet = crud.tra_cuu_luot_xe(db, schemas.TraCuuLuotXeParams(tu_ngay=tu_ngay, den_ngay=den_ngay, khu_vuc_id=khu_vuc_id))
    for i, l in enumerate(chi_tiet, start=2):
        gio_vao = l["thoi_gian_vao"].strftime("%d/%m/%Y %H:%M") if l["thoi_gian_vao"] else "-"
        gio_ra = l["thoi_gian_ra"].strftime("%d/%m/%Y %H:%M") if l["thoi_gian_ra"] else "-"
        hinh_thuc = l["hinh_thuc_gui"].value if hasattr(l["hinh_thuc_gui"], "value") else l["hinh_thuc_gui"]
        trang_thai = l["trang_thai"].value if hasattr(l["trang_thai"], "value") else l["trang_thai"]
        ws4.cell(row=i, column=1, value=l["bien_so"])
        ws4.cell(row=i, column=2, value=l["chu_xe"] or "-")
        ws4.cell(row=i, column=3, value=l["so_dien_thoai"] or "-")
        ws4.cell(row=i, column=4, value=l["ten_loai_xe"])
        ws4.cell(row=i, column=5, value=l["ten_khu_vuc"])
        ws4.cell(row=i, column=6, value=l["ma_vi_tri"])
        ws4.cell(row=i, column=7, value=NHAN_HINH_THUC_GUI.get(hinh_thuc, hinh_thuc or "-"))
        ws4.cell(row=i, column=8, value=gio_vao)
        ws4.cell(row=i, column=9, value=gio_ra)
        ws4.cell(row=i, column=10, value=l["phi_thu"] or 0)
        ws4.cell(row=i, column=11, value="Có" if l["mat_ve"] else "-")
        ws4.cell(row=i, column=12, value="Đang gửi" if trang_thai == "dang_gui" else "Hoàn tất")
    _tu_dong_gian_cot(ws4, {c: w for c, w in zip(
        [get_column_letter(i) for i in range(1, 13)],
        [14, 18, 14, 12, 20, 10, 24, 17, 17, 14, 9, 12],
    )})

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
